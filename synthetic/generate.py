"""
Generate synthetic yearly workbooks with the same layout as the real ones.

The real source is one Google Sheets workbook per year (2021-2026), exported
as XLSX. This module reproduces their structure closely enough that the
ingestion runs on the synthetic files unchanged, quirks included:

- **Two month-sheet layouts**: 2021 has no MacroCategoria column and its
  header sits on row 15; from 2022 on the header is on row 18.
- **Credits and debits side by side**: both sides share sheet rows, so most
  rows hold one credit *and* one debit.
- **Meaning in cell colors**: the credit amount's background encodes the
  owner (Person A / Person B), the debit's encodes the card. The same owner
  color appears either as an ``rgb`` value or as a ``theme`` index.
- **Installments**: card purchases split into ``(k/N)`` rows across later
  months, with the original purchase date, crossing year boundaries.
- **Cumulative portfolio sheet**: "Situação Atual" stacks one block per
  snapshot date, and each year's file repeats the whole history so far.
  Columns are added over time, and every block ends with a ``Total:`` row.
- **Dirty text**: case and whitespace drift, legacy categories missing from
  the canonical config, blank macro-categories, free text in value columns.

Output is deterministic for a given ``--seed`` and ``--as-of``.

CLI::

    python -m synthetic.generate --out data/synthetic
"""

from __future__ import annotations

import argparse
import calendar
import logging
import math
import random
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Color, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from synthetic import catalog as c

log = logging.getLogger("synthetic.generate")

FIRST_YEAR = 2021
LAST_YEAR = 2026
DEFAULT_SEED = 42
DEFAULT_AS_OF = date(2026, 10, 15)
DEFAULT_OUT = Path("data/synthetic")
FILE_TEMPLATE = "Planilha_de_Gastos_Financeiros {year}.xlsx"

CURRENCY_FORMAT = '"R$" #,##0.00'
DATE_FORMAT = "DD/MM/YYYY"
PCT_FORMAT = "0.0%"

#: Background color of a cell: ("rgb", "FF4BACC6"), ("theme", 9) or None.
Fill = tuple[str, str | int] | None
Month = tuple[int, int]

# Portfolio sheet: the first snapshot was written as free text, and columns
# were added to the table twice.
TEXT_SNAPSHOT_DATE = date(2022, 8, 25)
ALLOCATION_COLUMNS_SINCE = date(2023, 6, 1)
PROJECTION_COLUMNS_SINCE = date(2025, 1, 1)


# --------------------------------------------------------------------------
# Calendar helpers
# --------------------------------------------------------------------------


def add_months(year: int, month: int, k: int) -> Month:
    index = year * 12 + (month - 1) + k
    return index // 12, index % 12 + 1


def months_between(start: Month, end: Month) -> list[Month]:
    """Inclusive range of (year, month) pairs."""
    months = []
    current = start
    while current <= end:
        months.append(current)
        current = add_months(*current, 1)
    return months


def days_in(year: int, month: int) -> int:
    return calendar.monthrange(year, month)[1]


def last_business_day(year: int, month: int) -> date:
    day = date(year, month, days_in(year, month))
    while day.weekday() >= 5:
        day = day.replace(day=day.day - 1)
    return day


# --------------------------------------------------------------------------
# Ledger
# --------------------------------------------------------------------------


@dataclass
class Entry:
    """One side (credit or debit) of one sheet row."""

    description: str | None
    category: str | None
    amount: float
    when: date | None
    fill: Fill = None
    macro: str | None = None
    on_card: bool = False
    desc_fill: Fill = None
    """Debits only: the description cell's color, when it differs from the amount's."""
    status: str = "Sim"


@dataclass
class MonthBook:
    credits: list[Entry] = field(default_factory=list)
    debits: list[Entry] = field(default_factory=list)


@dataclass
class Holding:
    """One row of a portfolio snapshot block."""

    bank: str
    name: str
    value: float
    maturity: date | None
    mom: float
    yoy: float
    net_at_maturity: float | None
    return_to_date: float
    return_at_maturity: float | None


@dataclass
class Snapshot:
    when: date
    holdings: list[Holding]


@dataclass
class Asset:
    """A position simulated month by month to feed the snapshots."""

    bank: str
    name: str
    start: date
    initial: float
    monthly_return: float
    volatility: float = 0.0
    contribution: float = 0.0
    until: date | None = None
    """Last date the position is held (redeemed or matured)."""
    maturity: date | None = None
    """Maturity shown in the sheet; it can be earlier than `until`."""
    defaults_from: date | None = None
    """From this date on, the note loses value every month."""
    value: float = 0.0
    cost: float = 0.0
    history: dict[int, float] = field(default_factory=dict)


# --------------------------------------------------------------------------
# Generator
# --------------------------------------------------------------------------


class Generator:
    """Builds the ledger and the portfolio snapshots, then renders workbooks."""

    def __init__(self, seed: int = DEFAULT_SEED, as_of: date = DEFAULT_AS_OF) -> None:
        self.rng = random.Random(seed)
        self.as_of = as_of
        self.as_of_month: Month = (as_of.year, as_of.month)
        self.ledger: dict[Month, MonthBook] = defaultdict(MonthBook)
        self.snapshots: list[Snapshot] = []

    def build(self) -> Generator:
        self._card_purchases()
        self._subscriptions()
        for year, month in months_between((FIRST_YEAR, 1), self.as_of_month):
            self._income(year, month)
            self._bills(year, month)
        self._finalize()
        self._snapshots()
        return self

    # ---- helpers ---------------------------------------------------------

    def _in_range(self, month: Month) -> bool:
        return FIRST_YEAR <= month[0] <= LAST_YEAR

    def _inflation(self, year: int) -> float:
        return 1.055 ** (year - FIRST_YEAR)

    def _price(self, low: float, high: float, year: int) -> float:
        value = math.exp(self.rng.uniform(math.log(low), math.log(high)))
        return round(value * self._inflation(year), 2)

    def _day(self, year: int, month: int, low: int = 1, high: int = 28) -> date:
        return date(year, month, self.rng.randint(low, min(high, days_in(year, month))))

    def _person_a(self, year: int) -> Fill:
        if year == FIRST_YEAR:
            return None  # the 2021 workbook did not color credits by owner
        if self.rng.random() < 0.9:
            return ("rgb", c.PERSON_A_RGB)
        return ("theme", c.PERSON_A_THEME)

    def _person_b(self) -> Fill:
        if self.rng.random() < 0.85:
            return ("theme", c.PERSON_B_THEME)
        return ("rgb", c.PERSON_B_RGB)

    def _card(self, year: int) -> Fill:
        if year < c.CARD_C_SINCE_YEAR:
            cards, weights = [c.CARD_A_RGB, c.CARD_B_RGB], [75, 25]
        else:
            cards, weights = [c.CARD_A_RGB, c.CARD_B_RGB, c.CARD_C_RGB], [65, 22, 13]
        return ("rgb", self.rng.choices(cards, weights=weights)[0])

    def _messy(self, text: str) -> str:
        """Occasional case and whitespace drift, as typed by hand."""
        roll = self.rng.random()
        if roll < 0.03:
            return text[0] + text[1:].lower()
        if roll < 0.05 and " - " in text:
            return text.replace(" - ", " -  ", 1)
        return text

    def _macro(self, category: str) -> str | None:
        if self.rng.random() < 0.03:
            return None  # left blank by hand
        return c.EXPENSE_CATEGORIES.get(category, c.LEGACY_MACRO)

    # ---- debits ----------------------------------------------------------

    def _card_purchases(self) -> None:
        """Card purchases land on the next month's sheet, split in installments."""
        weights = [m.weight for m in c.MERCHANTS]
        for year, month in months_between((FIRST_YEAR - 1, 1), self.as_of_month):
            # The household grows, and so does the number of purchases.
            mean = 18 + 2 * (year - FIRST_YEAR + 1)
            for _ in range(max(8, round(self.rng.gauss(mean, 4)))):
                when = self._day(year, month, high=31)
                if when > self.as_of:
                    continue
                merchant = self.rng.choices(c.MERCHANTS, weights=weights)[0]
                self._post_purchase(merchant, when)

    def _post_purchase(self, merchant: c.Merchant, when: date) -> None:
        total = self._price(merchant.low, merchant.high, when.year)
        count = 1
        if merchant.installments and total > 250 and self.rng.random() < 0.7:
            count = self.rng.choice(merchant.installments)

        category = merchant.category
        legacy = c.LEGACY_CATEGORY_BY_MERCHANT.get(merchant.description)
        if legacy and when.year <= c.LEGACY_UNTIL_YEAR and self.rng.random() < 0.6:
            category = legacy

        fill = self._card(when.year)
        macro = self._macro(category)
        base = self._messy(merchant.description)
        item = self.rng.choice(merchant.items) if merchant.items else None
        tag_in_middle = self.rng.random() < 0.5
        amount = round(total / count, 2)

        for k in range(1, count + 1):
            month = add_months(when.year, when.month, k)
            if not self._in_range(month):
                continue
            if count == 1:
                desc = f"{base} - {item}" if item else base
            elif item is None:
                desc = f"{base} ({k}/{count})"
            elif tag_in_middle:
                desc = f"{base} ({k}/{count}) - {item}"
            else:
                desc = f"{base} - {item} ({k}/{count})"
            self.ledger[month].debits.append(
                Entry(desc, category, amount, when, fill, macro=macro, on_card=True)
            )

    def _subscriptions(self) -> None:
        for description, category, amount, since in c.SUBSCRIPTIONS:
            for year, month in months_between((since, 1), self.as_of_month):
                charged = add_months(year, month, -1)
                when = date(*charged, 7)
                self.ledger[(year, month)].debits.append(
                    Entry(
                        description,
                        category,
                        amount,
                        when,
                        ("rgb", c.CARD_A_RGB),
                        macro=c.EXPENSE_CATEGORIES[category],
                        on_card=True,
                    )
                )

    def _bills(self, year: int, month: int) -> None:
        """Bills paid by bank slip or transfer: no card, so no cell color."""
        book = self.ledger[(year, month)]
        yearly_adjustments = (year - FIRST_YEAR) + (1 if month >= 3 else 0)
        for description, category, base, jitter, since in c.BILLS:
            if year < since:
                continue
            if description == "Aluguel":
                amount = base * 1.05**yearly_adjustments
            else:
                amount = base * self._inflation(year) * (1 + self.rng.uniform(-jitter, jitter))
            book.debits.append(
                Entry(
                    description,
                    category,
                    round(amount, 2),
                    self._day(year, month, 5, 10),
                    macro=c.EXPENSE_CATEGORIES[category],
                )
            )

        if 2 <= month <= 11:
            book.debits.append(
                Entry(
                    f"IPTU ({month - 1}/10)",
                    "Imposto",
                    round(165 * self._inflation(year), 2),
                    self._day(year, month, 10, 15),
                    macro="Fixo",
                )
            )

        if self.rng.random() < 0.05:
            book.debits.append(
                Entry(
                    "Desconto antecipação fatura",
                    "Cartão de Crédito",
                    -round(self.rng.uniform(5, 40), 2),
                    self._day(year, month),
                    ("rgb", c.CARD_A_RGB),
                    macro="Variável",
                    on_card=True,
                )
            )

        income = sum(e.amount for e in book.credits)
        book.debits.append(
            Entry(
                "Aporte corretora",
                "Investimento",
                round(income * self.rng.uniform(0.12, 0.22), -1),
                self._day(year, month, 10, 20),
                macro="n/a",
            )
        )

    # ---- credits ---------------------------------------------------------

    def _income(self, year: int, month: int) -> None:
        book = self.ledger[(year, month)]
        raises = (year - FIRST_YEAR) + (1 if month >= 3 else 0)
        salary = 5200 * 1.07**raises
        inflation = self._inflation(year)

        def person_a(description: str, category: str, amount: float, when: date | None) -> None:
            book.credits.append(
                Entry(description, category, round(amount, 2), when, self._person_a(year))
            )

        def person_b(description: str, category: str, amount: float, when: date) -> None:
            book.credits.append(
                Entry(description, category, round(amount, 2), when, self._person_b())
            )

        person_a("Salário - Empresa A", "Salário", 0.4 * salary, date(year, month, 15))
        person_a("Salário - Empresa A", "Salário", 0.6 * salary, last_business_day(year, month))

        if year >= 2022:
            person_a("Vale-Alimentação", "Vale", 650 * inflation, self._day(year, month, 26, 29))
            if year < 2024:
                person_a("Vale-Refeição", "Vale", 350 * inflation, self._day(year, month, 26, 29))
            else:
                # Benefit merged into the food allowance; the row stayed, zeroed.
                person_a("Vale-Refeição", "Vale", 0, None)
            person_a("Auxílio Creche", "Auxílio", 480 * inflation, last_business_day(year, month))
        if year >= 2024:
            person_a("Auxílio Celular", "Auxílio", 60, date(year, month, 15))

        if month == 1:
            person_a("Participação de Resultados", "PLR", 2.4 * salary, date(year, 1, 30))
        if month == 7:
            person_a("Participação de Resultados", "PLR", 1.1 * salary, date(year, 7, 30))
        if month == 11:
            person_a("Décimo Terceiro", "Décimo Terceiro", 0.5 * salary, date(year, 11, 30))
        if month == 12:
            person_a("Décimo Terceiro", "Décimo Terceiro", 0.5 * salary, date(year, 12, 20))
        if month == 1 + (year * 7) % 12:
            person_a("Férias", "Férias", 1.33 * salary, self._day(year, month, 1, 5))

        if year >= 2022:
            person_b(
                "Pró-labore - Empresa B",
                "Empresa",
                2400 * 1.08 ** (year - 2022),
                date(year, month, 25),
            )
            if self.rng.random() < 0.12:
                person_b(
                    "Freelance - design",
                    "Freelance",
                    self.rng.uniform(400, 1500),
                    self._day(year, month),
                )

        if year >= 2022 and self.rng.random() < 0.15:
            person_a("Cashback cartão", "Cashback", self.rng.uniform(5, 60), self._day(year, month))

        if (year, month) >= (2022, 9):
            months_invested = len(months_between((2022, 9), (year, month)))
            tickers = c.FUND_TICKERS[: min(len(c.FUND_TICKERS), 2 + months_invested // 3)]
            dividends = [
                Entry(
                    ticker,
                    "Proventos",
                    round(self.rng.uniform(3, 45) * (1 + months_invested / 24), 2),
                    self._day(year, month, 8, 20),
                    ("rgb", c.NEUTRAL_RGB),
                )
                for ticker in tickers
            ]
            book.credits.extend(sorted(dividends, key=lambda e: e.when or date.max))

    # ---- status and finishing touches -----------------------------------

    def _status(self, entry: Entry, month: Month) -> str:
        if month < self.as_of_month:
            return "Não" if self.rng.random() < 0.01 else "Sim"
        if month > self.as_of_month:
            return "Não"
        if entry.on_card:  # the card bill is due on the 10th
            return "Sim" if self.as_of.day >= 10 else "Não"
        return "Sim" if entry.when is not None and entry.when <= self.as_of else "Não"

    def _finalize(self) -> None:
        for month in sorted(self.ledger):
            book = self.ledger[month]
            book.debits.sort(key=lambda e: e.when or date.max)
            for entry in book.credits:
                entry.status = self._status(entry, month)
            for entry in book.debits:
                entry.status = self._status(entry, month)
                if month <= self.as_of_month and self.rng.random() < 0.01:
                    entry.desc_fill = ("rgb", c.STRAY_HIGHLIGHT_RGB)

    # ---- portfolio -------------------------------------------------------

    def _portfolio_assets(self) -> list[Asset]:
        rng = self.rng
        start = TEXT_SNAPSHOT_DATE
        assets = [
            Asset("Banco Público", "FGTS", start, 3500, 0.003, contribution=300),
            Asset("Corretora Beta", "FIIs", start, 4200, 0.004, 0.03, contribution=650),
            Asset(
                "Banco Alfa (Fundos)",
                "Fundo Multimercado Alfa",
                start,
                8000,
                0.008,
                0.01,
                until=date(2025, 6, 30),
            ),
            Asset(
                "Banco Alfa (Fundos)",
                "Fundo Renda Fixa Alfa",
                start,
                7500,
                0.009,
                until=date(2024, 9, 30),
            ),
            Asset(
                "Banco Alfa (Fundos)",
                "Fundo Long Short Alfa",
                start,
                9000,
                0.007,
                0.015,
                until=date(2025, 12, 31),
            ),
            Asset(
                "Banco Alfa (Renda Fixa)",
                "CDB Banco Ômega IPCA+6,5%",
                start,
                7000,
                0.0095,
                until=date(2026, 4, 15),
                maturity=date(2026, 4, 15),
            ),
            Asset(
                "Banco Alfa (Renda Fixa)",
                "LCI Financeira Sigma 12,5% a.a.",
                start,
                3800,
                0.0098,
                until=date(2024, 5, 15),
                maturity=date(2024, 5, 15),
            ),
            Asset(
                "Fintech Delta",
                "CDB 120% CDI Liq. Diária",
                start,
                3000,
                0.011,
                until=date(2022, 12, 12),
                maturity=date(2022, 12, 12),
            ),
            Asset(
                "Fintech Delta",
                "CDB 110% CDI Liq. Diária",
                start,
                2500,
                0.01,
                until=date(2023, 6, 30),
                maturity=date(2023, 6, 30),
            ),
            Asset("Fintech Delta", "Saldo (CDI 100%)", start, 2000, 0.009, 0.0, contribution=150),
            Asset(
                "Fintech Delta",
                "Tesouro IPCA+ 2035",
                date(2024, 3, 1),
                3000,
                0.008,
                contribution=400,
                maturity=date(2035, 5, 15),
            ),
            Asset(
                "Fintech Delta",
                "Tesouro IPCA+ 2045",
                date(2025, 2, 1),
                5000,
                0.008,
                contribution=300,
                maturity=date(2045, 5, 15),
            ),
            Asset(
                "Plataforma Kappa", "Startup Equity", start, 3000, 0.0, maturity=date(2026, 6, 30)
            ),
            Asset("Carteira Digital Zeta", "180% CDI", start, 300, 0.012),
            Asset("Corretora Cripto", "Bitcoin", date(2024, 6, 1), 800, 0.02, 0.12),
            Asset("Previdência Épsilon", "PGBL", date(2025, 9, 1), 3000, 0.008, contribution=500),
        ]
        projects = rng.sample(c.NOTE_PROJECTS, len(c.NOTE_PROJECTS))
        for k, project in enumerate(projects):
            year, month = add_months(2023, 3, 2 * k)
            note_start = date(year, month, 14)
            if note_start > self.as_of:
                break
            defaults = rng.random() < 0.25
            assets.append(
                Asset(
                    "Plataforma Gama",
                    f"{rng.choice(c.NOTE_PREFIXES)} {project}",
                    note_start,
                    round(rng.uniform(800, 3000), -1),
                    0.009,
                    maturity=date(
                        rng.randint(2027, 2030), rng.randint(1, 12), rng.choice([7, 14, 21, 28])
                    ),
                    defaults_from=date(*add_months(year, month, rng.randint(10, 24)), 1)
                    if defaults
                    else None,
                )
            )
        return assets

    def _snapshot_dates(self) -> list[date]:
        dates = [TEXT_SNAPSHOT_DATE, date(2022, 9, 26)]
        for year, month in months_between((2022, 10), self.as_of_month):
            when = self._day(year, month, 20, 28)
            if when <= self.as_of:
                dates.append(when)
        return dates

    def _display_bank(self, bank: str) -> str:
        roll = self.rng.random()
        if roll < 0.03:
            return bank + " "
        if roll < 0.05:
            return bank[0] + bank[1:].lower()
        return bank

    def _snapshots(self) -> None:
        assets = self._portfolio_assets()
        for index, when in enumerate(self._snapshot_dates()):
            holdings = []
            for asset in assets:
                if asset.start > when or (asset.until is not None and when > asset.until):
                    continue
                if not asset.history:
                    asset.value = asset.cost = asset.initial
                    previous = None
                else:
                    previous = asset.value
                    rate = asset.monthly_return
                    if asset.defaults_from is not None and when >= asset.defaults_from:
                        rate = -0.08
                    growth = 1 + self.rng.gauss(rate, asset.volatility)
                    asset.value = max(0.0, asset.value * growth + asset.contribution)
                    asset.cost += asset.contribution
                asset.history[index] = asset.value

                year_ago = asset.history.get(index - 12)
                net = None
                if asset.maturity is not None and asset.maturity > when:
                    months_left = (asset.maturity - when).days / 30.4
                    net = asset.value * (1 + asset.monthly_return) ** months_left * 0.85
                name = asset.name
                if asset.bank == "Carteira Digital Zeta" and when >= date(2026, 1, 1):
                    name = "Alterado para 110% CDI"  # a note typed over the asset name

                holdings.append(
                    Holding(
                        bank=self._display_bank(asset.bank),
                        name=name,
                        value=round(asset.value, 2),
                        maturity=asset.maturity,
                        mom=asset.value / previous - 1 if previous else 0.0,
                        yoy=asset.value / year_ago - 1 if year_ago else 0.0,
                        net_at_maturity=round(net, 2) if net is not None else None,
                        return_to_date=asset.value / asset.cost - 1,
                        return_at_maturity=net / asset.cost - 1 if net is not None else None,
                    )
                )
            self.snapshots.append(Snapshot(when, holdings))


# --------------------------------------------------------------------------
# Rendering
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class SheetLayout:
    header_row: int
    title_cell: str
    debit_cols: dict[str, int]
    debit_computed_col: int


CREDIT_COLS = {"description": 2, "category": 3, "amount": 4, "date": 5, "status": 6}
CREDIT_COMPUTED_COL = 7

LAYOUT_2021 = SheetLayout(
    header_row=15,
    title_cell="C4",
    debit_cols={"description": 9, "category": 10, "amount": 11, "date": 12, "status": 13},
    debit_computed_col=14,
)
LAYOUT_2022 = SheetLayout(
    header_row=18,
    title_cell="C6",
    debit_cols={
        "description": 9,
        "macro": 10,
        "category": 11,
        "amount": 12,
        "date": 13,
        "status": 14,
    },
    debit_computed_col=15,
)

HEADER_LABELS = {
    "description": "Descrição",
    "macro": "MacroCategoria",
    "category": "Categoria",
    "amount": " Valor",
    "date": "Data",
    "status": "OK",
}


def _fill(cell, fill: Fill) -> None:
    if fill is None:
        return
    kind, value = fill
    color = Color(rgb=value) if kind == "rgb" else Color(theme=value, tint=0.0)
    cell.fill = PatternFill(fill_type="solid", start_color=color, end_color=color)


def _money(ws: Worksheet, row: int, col: int, value: float) -> None:
    cell = ws.cell(row, col, round(value, 2))
    cell.number_format = CURRENCY_FORMAT


def _write_entry(ws: Worksheet, row: int, cols: dict[str, int], entry: Entry, debit: bool) -> None:
    desc = ws.cell(row, cols["description"], entry.description)
    ws.cell(row, cols["category"], entry.category)
    if "macro" in cols:
        ws.cell(row, cols["macro"], entry.macro)
    amount = ws.cell(row, cols["amount"], entry.amount)
    amount.number_format = CURRENCY_FORMAT
    if entry.when is not None:
        ws.cell(row, cols["date"], entry.when).number_format = DATE_FORMAT
    ws.cell(row, cols["status"], entry.status)

    _fill(amount, entry.fill)
    if debit:
        _fill(desc, entry.desc_fill or entry.fill)


def _write_summary(ws: Worksheet, year: int, book: MonthBook) -> None:
    credits = sum(e.amount for e in book.credits)
    debits = sum(e.amount for e in book.debits)
    by_card: dict[str, float] = defaultdict(float)
    for entry in book.debits:
        if entry.fill is not None:
            by_card[str(entry.fill[1])] += entry.amount
    balance = credits - debits
    goal = 0.2 * credits

    ws["B2"] = "Planilha de Gastos / Orçamento Financeiro"
    ws["I2"] = "Resumo dos Créditos e Débitos"
    if year == FIRST_YEAR:
        ws["I3"], ws["J3"] = "Créditos:", credits
        ws["I4"], ws["J4"], ws["K4"] = "Débitos:", "Card A:", by_card[c.CARD_A_RGB]
        ws["I5"], ws["J5"], ws["K5"] = "Débitos:", "Card B:", by_card[c.CARD_B_RGB]
        ws["I6"], ws["J6"] = "Débitos:", debits
        ws["I7"], ws["J7"] = "Saldo:", balance
        ws["I9"], ws["J9"] = "Meta  do mês:", goal
        ws["I10"], ws["K10"] = "Meta Atingida", balance - goal
        return

    ws["I3"], ws["K3"] = "Créditos:", credits
    ws["K4"] = " Person A:"
    ws["K5"] = " Person B:"
    ws["I6"], ws["K6"] = "Débitos:", debits
    ws["I7"], ws["K7"], ws["L7"] = "Cartão de crédito:", "Card A:", by_card[c.CARD_A_RGB]
    ws["K8"], ws["L8"] = "Card B:", by_card[c.CARD_B_RGB]
    ws["K9"], ws["L9"] = " Card C:", by_card[c.CARD_C_RGB]
    ws["I10"], ws["K10"] = "Saldo:", balance
    ws["I12"], ws["K12"] = "Meta do mês:", goal
    ws["I13"], ws["L13"] = "Meta Atingida", balance - goal
    ws["B8"], ws["D8"] = "Previsto", " Gasto"
    for row, (share, label) in enumerate(
        [
            (0.3, "30% Lazer e gastos variáveis"),
            (0.5, "50% Gastos fixos"),
            (0.1, "10% investimentos futuros"),
            (0.1, "10% Reserva emergência"),
        ],
        start=9,
    ):
        _money(ws, row, 2, share * credits)
        ws.cell(row, 3, label)


def render_month(ws: Worksheet, year: int, month: int, book: MonthBook) -> None:
    layout = LAYOUT_2021 if year == FIRST_YEAR else LAYOUT_2022
    _write_summary(ws, year, book)

    title = f"{c.MONTH_NAMES[month - 1]} {year}"
    ws[layout.title_cell] = title if year == FIRST_YEAR else title.upper()

    ws.cell(layout.header_row - 2, 2, "Créditos (Entradas)")
    ws.cell(layout.header_row - 2, 9, "Débitos (Saídas)")
    for name, col in CREDIT_COLS.items():
        ws.cell(layout.header_row, col, HEADER_LABELS[name])
    for name, col in layout.debit_cols.items():
        ws.cell(layout.header_row, col, HEADER_LABELS[name])

    first = layout.header_row + 1
    for i in range(max(len(book.credits), len(book.debits))):
        row = first + i
        credit = book.credits[i] if i < len(book.credits) else None
        debit = book.debits[i] if i < len(book.debits) else None
        # The computed column is a formula filled down the whole credit side.
        paid = credit.amount if credit and credit.status == "Sim" else 0
        _money(ws, row, CREDIT_COMPUTED_COL, paid)
        if credit is not None:
            _write_entry(ws, row, CREDIT_COLS, credit, debit=False)
        if debit is not None:
            _write_entry(ws, row, layout.debit_cols, debit, debit=True)
            paid = debit.amount if debit.status == "Sim" else 0
            _money(ws, row, layout.debit_computed_col, paid)


def _portfolio_headers(when: date) -> dict[int, str]:
    headers = {2: "Banco/Corretora", 3: "Aplicação", 4: "Valor Bruto", 6: "Vencimento"}
    if when >= ALLOCATION_COLUMNS_SINCE:
        headers.update({5: "Alocação", 7: "MoM", 8: "YoY"})
    if when >= PROJECTION_COLUMNS_SINCE:
        headers.update({9: "Valor líquido final", 10: "Rendimento atual", 11: "Rendimento final"})
    return headers


def render_portfolio(ws: Worksheet, snapshots: list[Snapshot]) -> None:
    row = 2
    for snap in snapshots:
        ws.cell(row, 2, "Data Referência:")
        ws.cell(row, 3, snap.when).number_format = DATE_FORMAT
        row += 1
        total = sum(h.value for h in snap.holdings)

        if snap.when == TEXT_SNAPSHOT_DATE:
            # The first snapshot was a plain-text summary, not a table.
            by_bank: dict[str, float] = defaultdict(float)
            for holding in snap.holdings:
                by_bank[holding.bank.strip()] += holding.value
            for bank, value in by_bank.items():
                ws.cell(row, 2, f"{value / 1000:.1f}k em {bank}")
                row += 1
            ws.cell(row, 2, f"{total / 1000:.1f}k Total")
            row += 2
            continue

        headers = _portfolio_headers(snap.when)
        for col, label in headers.items():
            ws.cell(row, col, label)
        row += 1
        for holding in snap.holdings:
            ws.cell(row, 2, holding.bank)
            ws.cell(row, 3, holding.name)
            _money(ws, row, 4, holding.value)
            if holding.maturity is not None:
                ws.cell(row, 6, holding.maturity).number_format = DATE_FORMAT
            else:
                ws.cell(row, 6, "N/A")
            if 5 in headers:
                ws.cell(row, 5, holding.value / total).number_format = PCT_FORMAT
                ws.cell(row, 7, holding.mom).number_format = PCT_FORMAT
                ws.cell(row, 8, holding.yoy).number_format = PCT_FORMAT
            if 9 in headers:
                if holding.net_at_maturity is not None:
                    _money(ws, row, 9, holding.net_at_maturity)
                ws.cell(row, 10, holding.return_to_date).number_format = PCT_FORMAT
                if holding.return_at_maturity is not None:
                    ws.cell(row, 11, holding.return_at_maturity).number_format = PCT_FORMAT
            row += 1

        # Blank line, then the total, with its label in column C, not B.
        row += 1
        ws.cell(row, 3, "Total:")
        _money(ws, row, 4, total)
        if 7 in headers:
            weights = [h.value / total for h in snap.holdings]
            mom = sum(w * h.mom for w, h in zip(weights, snap.holdings, strict=True))
            yoy = sum(w * h.yoy for w, h in zip(weights, snap.holdings, strict=True))
            ws.cell(row, 7, mom).number_format = PCT_FORMAT
            ws.cell(row, 8, yoy).number_format = PCT_FORMAT
        row += 2


def render_config(ws: Worksheet) -> None:
    ws["I2"] = "Descrição das Categorias:"
    ws["I3"], ws["J3"], ws["K3"] = "Receitas", "Despesas", "MacroCategoria"
    expenses = list(c.EXPENSE_CATEGORIES.items())
    for i, row in enumerate(range(4, 39)):
        ws.cell(row, 9, c.INCOME_CATEGORIES[i] if i < len(c.INCOME_CATEGORIES) else "-")
        if i < len(expenses):
            ws.cell(row, 10, expenses[i][0])
            ws.cell(row, 11, expenses[i][1])
        else:
            ws.cell(row, 10, "-")


def render_reports(ws: Worksheet, year: int, ledger: dict[Month, MonthBook]) -> None:
    ws["B2"] = f"Relatório anual {year}"
    ws["B4"], ws["C4"], ws["D4"] = "Mês", "Créditos", "Débitos"
    for month in range(1, 13):
        book = ledger.get((year, month), MonthBook())
        ws.cell(4 + month, 2, c.MONTH_NAMES[month - 1])
        _money(ws, 4 + month, 3, sum(e.amount for e in book.credits))
        _money(ws, 4 + month, 4, sum(e.amount for e in book.debits))


def write_workbooks(generator: Generator, out_dir: Path) -> list[Path]:
    """Render one workbook per year into `out_dir`; returns the file paths."""
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for year in range(FIRST_YEAR, LAST_YEAR + 1):
        wb = Workbook()
        wb.remove(wb.active)
        for month in range(1, 13):
            book = generator.ledger.get((year, month), MonthBook())
            render_month(wb.create_sheet(c.MONTH_SHEETS[month - 1]), year, month, book)
        render_reports(wb.create_sheet("Relatórios"), year, generator.ledger)
        if year >= TEXT_SNAPSHOT_DATE.year:
            history = [s for s in generator.snapshots if s.when.year <= year]
            render_portfolio(wb.create_sheet("Situação Atual"), history)
        render_config(wb.create_sheet("Configurações"))

        path = out_dir / FILE_TEMPLATE.format(year=year)
        wb.save(path)
        paths.append(path)
    return paths


def generate(
    out_dir: str | Path = DEFAULT_OUT, seed: int = DEFAULT_SEED, as_of: date = DEFAULT_AS_OF
) -> list[Path]:
    generator = Generator(seed=seed, as_of=as_of).build()
    paths = write_workbooks(generator, Path(out_dir))
    rows = sum(len(b.credits) + len(b.debits) for b in generator.ledger.values())
    log.info(
        "Wrote %d workbooks to %s (%d transactions, %d portfolio snapshots)",
        len(paths),
        out_dir,
        rows,
        len(generator.snapshots),
    )
    return paths


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        prog="python -m synthetic.generate",
        description="Generate synthetic yearly workbooks with the real layout.",
    )
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="Output directory.")
    parser.add_argument("--seed", type=int, default=DEFAULT_SEED, help="Random seed.")
    parser.add_argument(
        "--as-of",
        type=date.fromisoformat,
        default=DEFAULT_AS_OF,
        help="Last day with data, YYYY-MM-DD (later months hold only future installments).",
    )
    args = parser.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    generate(args.out, args.seed, args.as_of)


if __name__ == "__main__":
    main()
