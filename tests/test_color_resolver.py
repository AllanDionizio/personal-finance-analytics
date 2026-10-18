"""Testes do color_resolver.

Inclui testes puros pra `dono_from_hex` / `meio_pagamento_from_hex` e
testes de integração sobre o XLSX sintético (ou um real, via env var).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from ingestion.color_resolver import (
    DONO_BY_HEX,
    DONO_DEFAULT,
    MEIO_PAGAMENTO_BY_HEX,
    MEIO_PAGAMENTO_DEFAULT,
    dono_from_hex,
    load_theme_palette,
    meio_pagamento_from_hex,
    resolve_cell_color,
    resolve_dono,
    resolve_meio_pagamento,
)

# --------------------------------------------------------------------------
# Testes puros
# --------------------------------------------------------------------------


class TestDonoFromHex:
    """Mapeamento de hex pra dono."""

    @pytest.mark.parametrize(
        "hex_color,esperado",
        [
            ("4BACC6", "Person A"),
            ("F79646", "Person B"),
            ("FFFFFF", "Outros"),
            ("000000", "Outros"),
        ],
    )
    def test_cores_conhecidas(self, hex_color: str, esperado: str) -> None:
        assert dono_from_hex(hex_color) == esperado

    def test_case_insensitive(self) -> None:
        assert dono_from_hex("4bacc6") == "Person A"
        assert dono_from_hex("4BaCc6") == "Person A"

    def test_com_prefixo_hash(self) -> None:
        assert dono_from_hex("#4BACC6") == "Person A"
        assert dono_from_hex("#f79646") == "Person B"

    def test_cor_desconhecida_retorna_default(self) -> None:
        assert dono_from_hex("ABCDEF") == DONO_DEFAULT

    def test_none_retorna_default(self) -> None:
        assert dono_from_hex(None) == DONO_DEFAULT

    def test_mapeamento_canonico_contem_dimensoes_esperadas(self) -> None:
        """Sanity check: os valores únicos do mapeamento são os esperados."""
        assert set(DONO_BY_HEX.values()) == {"Person A", "Person B", "Outros"}


class TestMeioPagamentoFromHex:
    """Mapeamento de hex pra meio de pagamento."""

    @pytest.mark.parametrize(
        "hex_color,esperado",
        [
            ("B4A7D6", "Card A"),
            ("8E7CC3", "Card B"),
            ("b4a7d6", "Card A"),
            ("#8e7cc3", "Card B"),
            ("9FC5E8", "Card C"),
        ],
    )
    def test_cores_conhecidas(self, hex_color: str, esperado: str) -> None:
        assert meio_pagamento_from_hex(hex_color) == esperado

    def test_cor_desconhecida_retorna_default(self) -> None:
        assert meio_pagamento_from_hex("ABCDEF") == MEIO_PAGAMENTO_DEFAULT

    def test_none_retorna_default(self) -> None:
        assert meio_pagamento_from_hex(None) == MEIO_PAGAMENTO_DEFAULT

    def test_mapeamento_canonico(self) -> None:
        assert set(MEIO_PAGAMENTO_BY_HEX.values()) == {"Card A", "Card B", "Card C"}


# --------------------------------------------------------------------------
# Testes de integração — XLSX sintético ou real
# --------------------------------------------------------------------------


class TestThemePalette:
    """Lê a paleta do tema do XLSX."""

    def test_paleta_tem_12_entradas(self, fixture_xlsx_path: Path) -> None:
        palette = load_theme_palette(fixture_xlsx_path)
        # OOXML padrão tem 12 cores no clrScheme
        assert len(palette) == 12

    def test_paleta_contem_accent5_person_a(self, fixture_xlsx_path: Path) -> None:
        """Accent5 (índice 8) deve ser o teal usado pra Person A."""
        palette = load_theme_palette(fixture_xlsx_path)
        assert palette[8] == "4BACC6"

    def test_paleta_contem_accent6_person_b(self, fixture_xlsx_path: Path) -> None:
        """Accent6 (índice 9) deve ser o laranja usado pra Person B."""
        palette = load_theme_palette(fixture_xlsx_path)
        assert palette[9] == "F79646"

    def test_paleta_arquivo_inexistente(self) -> None:
        with pytest.raises(FileNotFoundError):
            load_theme_palette("/tmp/nao_existe_zzz.xlsx")


@pytest.fixture(scope="module")
def wb(fixture_xlsx_path: Path):
    from openpyxl import load_workbook

    return load_workbook(fixture_xlsx_path, data_only=True)


@pytest.fixture(scope="module")
def palette(fixture_xlsx_path: Path) -> list[str | None]:
    return load_theme_palette(fixture_xlsx_path)


class TestResolveCellColorReal:
    """Resolve cor em células do XLSX usando openpyxl."""

    MESES = ("JAN", "FEV", "MAR", "ABR", "MAI", "JUN", "JUL", "AGO", "SET", "OUT", "NOV", "DEZ")

    def _primeira_celula(self, wb, col: int, source_type: str, hex_color: str, palette):
        """Primeira célula da coluna, em qualquer aba mensal, cuja cor vem de
        `source_type` e resolve para `hex_color`."""
        for sheet in self.MESES:
            ws = wb[sheet]
            for row in range(16, ws.max_row + 1):
                cell = ws.cell(row, col)
                if cell.value is None:
                    continue
                result = resolve_cell_color(cell, palette)
                if result.source_type == source_type and result.hex == hex_color:
                    return cell
        pytest.fail(f"Nenhuma célula {source_type} {hex_color} na coluna {col}")

    def test_resolve_cor_rgb_direto(self, wb, palette) -> None:
        """Entrada de Person A gravada como rgb FF4BACC6 → 4BACC6."""
        cell = self._primeira_celula(wb, 4, "rgb", "4BACC6", palette)
        assert resolve_dono(cell, palette) == "Person A"

    def test_resolve_cor_theme(self, wb, palette) -> None:
        """Entrada de Person B gravada como theme=9 → F79646."""
        cell = self._primeira_celula(wb, 4, "theme", "F79646", palette)
        assert resolve_dono(cell, palette) == "Person B"

    def test_mesma_cor_em_rgb_e_theme_resolve_igual(self, wb, palette) -> None:
        """O quirk do Google Sheets: a cor de Person A aparece nas duas formas."""
        rgb = self._primeira_celula(wb, 4, "rgb", "4BACC6", palette)
        theme = self._primeira_celula(wb, 4, "theme", "4BACC6", palette)
        assert resolve_dono(rgb, palette) == resolve_dono(theme, palette) == "Person A"

    def test_resolve_card_a_de_saida(self, wb, palette) -> None:
        """Saídas (col L) com cor B4A7D6 devem mapear como Card A."""
        cell = self._primeira_celula(wb, 12, "rgb", "B4A7D6", palette)
        assert resolve_meio_pagamento(cell, palette) == "Card A"
