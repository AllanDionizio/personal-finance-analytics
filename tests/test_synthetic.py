"""Tests for the synthetic workbook generator."""

from __future__ import annotations

import itertools
import re
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from ingestion.extract_sheets import detect_month_layout
from ingestion.snapshot_blocks import find_snapshot_blocks
from synthetic import catalog
from synthetic.generate import (
    ASSET_SALE,
    DEFAULT_AS_OF,
    FILE_TEMPLATE,
    FIRST_YEAR,
    LAST_YEAR,
    TEXT_SNAPSHOT_DATE,
    Generator,
)

INSTALLMENT_RE = re.compile(r"\(\d+/\d+\)")


def _fingerprint(generator: Generator) -> list[tuple]:
    return [
        (month, side, e.description, e.category, e.amount, e.when, e.fill, e.status)
        for month, book in sorted(generator.ledger.items())
        for side, entries in (("credit", book.credits), ("debit", book.debits))
        for e in entries
    ]


@pytest.fixture(scope="module")
def generator() -> Generator:
    return Generator().build()


def _workbook(synthetic_dir: Path, year: int):
    return load_workbook(synthetic_dir / FILE_TEMPLATE.format(year=year), data_only=True)


def test_same_seed_gives_same_data(generator: Generator) -> None:
    assert _fingerprint(Generator().build()) == _fingerprint(generator)


def test_different_seed_gives_different_data(generator: Generator) -> None:
    assert _fingerprint(Generator(seed=7).build()) != _fingerprint(generator)


def test_one_workbook_per_year(synthetic_dir: Path) -> None:
    names = sorted(p.name for p in synthetic_dir.glob("*.xlsx"))
    assert names == [FILE_TEMPLATE.format(year=y) for y in range(FIRST_YEAR, LAST_YEAR + 1)]


def test_2021_layout_has_no_macro_category(synthetic_dir: Path) -> None:
    layout = detect_month_layout(_workbook(synthetic_dir, 2021)["JAN"])
    assert layout is not None
    assert layout.header_row == 15
    assert "macro_category" not in layout.debit


def test_later_layout_has_macro_category(synthetic_dir: Path) -> None:
    layout = detect_month_layout(_workbook(synthetic_dir, 2026)["JAN"])
    assert layout is not None
    assert layout.header_row == 18
    assert "macro_category" in layout.debit


@pytest.mark.parametrize("as_of", [date(FIRST_YEAR - 1, 12, 31), date(LAST_YEAR + 1, 1, 1)])
def test_as_of_outside_workbook_years_is_rejected(as_of: date) -> None:
    with pytest.raises(ValueError, match="as_of"):
        Generator(as_of=as_of)


def test_installments_carry_over_from_previous_year(generator: Generator) -> None:
    january = generator.ledger[(FIRST_YEAR, 1)].debits
    carried = [e for e in january if e.when and e.when.year < FIRST_YEAR]
    assert all(e.on_card for e in carried)
    assert any(INSTALLMENT_RE.search(e.description or "") for e in carried)


def test_months_after_as_of_hold_only_unpaid_card_charges(generator: Generator) -> None:
    future = [book for month, book in generator.ledger.items() if month > generator.as_of_month]
    assert future
    for book in future:
        assert book.credits == []
        assert all(e.on_card and e.when <= generator.as_of for e in book.debits)
        assert all(e.status == "Não" for e in book.debits)
    assert any(INSTALLMENT_RE.search(e.description or "") for e in future[-1].debits)


def test_owner_colors_use_both_rgb_and_theme(generator: Generator) -> None:
    fills = {e.fill for book in generator.ledger.values() for e in book.credits}
    assert ("rgb", catalog.PERSON_A_RGB) in fills
    assert ("theme", catalog.PERSON_A_THEME) in fills
    assert ("rgb", catalog.PERSON_B_RGB) in fills
    assert ("theme", catalog.PERSON_B_THEME) in fills


def test_legacy_categories_only_in_old_years(generator: Generator) -> None:
    legacy = set(catalog.LEGACY_CATEGORY_BY_MERCHANT.values())
    years = {
        e.when.year
        for book in generator.ledger.values()
        for e in book.debits
        if e.category in legacy and e.when
    }
    assert years
    assert max(years) <= catalog.LEGACY_UNTIL_YEAR


def test_portfolio_sheet_is_cumulative(synthetic_dir: Path) -> None:
    for year in range(2022, LAST_YEAR + 1):
        blocks = find_snapshot_blocks(_workbook(synthetic_dir, year)["Situação Atual"])
        dates = [b.reference_date for b in blocks]
        assert dates[0] == date(2022, 8, 25)
        assert max(dates) <= min(date(year, 12, 31), DEFAULT_AS_OF)


def test_no_portfolio_sheet_before_2022(synthetic_dir: Path) -> None:
    assert "Situação Atual" not in _workbook(synthetic_dir, 2021).sheetnames


def _table_snapshots(generator: Generator) -> list:
    """Snapshots loaded as tables (the first one was free text)."""
    return [s for s in generator.snapshots if s.when != TEXT_SNAPSHOT_DATE]


def test_each_contribution_falls_between_two_snapshots(generator: Generator) -> None:
    """The budget books one contribution per portfolio period, inside it."""
    snapshots = _table_snapshots(generator)
    for previous, current in itertools.pairwise(snapshots):
        month = generator.ledger[(current.when.year, current.when.month)]
        contributions = [e for e in month.debits if e.category == "Investimento"]
        assert len(contributions) == 1, current.when
        assert previous.when < contributions[0].when <= current.when


def test_maturing_positions_stay_in_the_portfolio(generator: Generator) -> None:
    """A position that ends is reinvested, so the total never falls off a cliff.

    Before reinvestment, a matured position vanished with no money booked
    anywhere, and the implied return of that period was a large loss.
    """
    totals = [sum(h.value for h in s.holdings) for s in _table_snapshots(generator)]
    for previous, current in itertools.pairwise(totals):
        assert current >= 0.95 * previous


def test_asset_sale_is_typed_on_the_income_side(generator: Generator) -> None:
    sale = [
        e
        for book in generator.ledger.values()
        for e in book.credits
        if e.category == "Investimentos"
    ]
    assert len(sale) == ASSET_SALE[3]
    assert all(INSTALLMENT_RE.search(e.description) for e in sale)
